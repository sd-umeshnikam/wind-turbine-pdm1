"""Shared visual identity and helper builders for the platform's Word deliverables.

Palette/fonts intentionally match the project's established PdM case-study identity
(light industrial: copper/amber accent, teal/amber/rust semantic colors, Bahnschrift
headings, Georgia body, Consolas mono) - Windows-native fonts chosen so these render
correctly without font installation.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Inches

REPO_ROOT = Path(__file__).resolve().parents[3]
CHARTS_DIR = REPO_ROOT / "eda" / "outputs" / "charts"

HEADING_FONT = "Bahnschrift"
BODY_FONT = "Georgia"
MONO_FONT = "Consolas"

# Semantic palette
COPPER = RGBColor(0xB8, 0x73, 0x33)      # primary accent
DARK_TEAL = RGBColor(0x1F, 0x3B, 0x3A)   # headings / dark text
SLATE = RGBColor(0x3A, 0x44, 0x47)       # secondary body text
TEAL = RGBColor(0x2F, 0x6F, 0x62)        # "normal / good"
AMBER = RGBColor(0xC9, 0x7B, 0x2B)       # "warning"
RUST = RGBColor(0xA6, 0x3A, 0x2E)        # "critical"
LIGHT_FILL = "F2EFE9"                    # table header fill (hex, no #)
CODE_FILL = "EDEDE8"


def set_cell_shading(cell, hex_color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def set_paragraph_shading(paragraph, hex_color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    paragraph.paragraph_format.element.get_or_add_pPr().append(shd)


def add_left_border(paragraph, color_rgb: RGBColor, size: int = 24) -> None:
    """Colored left rule on a paragraph, used for callout boxes."""
    pPr = paragraph.paragraph_format.element.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), "%02X%02X%02X" % (color_rgb[0], color_rgb[1], color_rgb[2]))
    pBdr.append(left)
    pPr.append(pBdr)


def new_document() -> Document:
    doc = Document()

    section = doc.sections[0]
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = DARK_TEAL
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for level, size, color in [(1, 20, COPPER), (2, 15, DARK_TEAL), (3, 12.5, DARK_TEAL)]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = HEADING_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(18 if level == 1 else 12)
        style.paragraph_format.space_after = Pt(8)

    return doc


def add_title_page(doc: Document, title: str, subtitle: str, meta_lines: list[str]) -> None:
    for _ in range(3):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.font.name = HEADING_FONT
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = COPPER

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitle)
    run.font.name = HEADING_FONT
    run.font.size = Pt(15)
    run.font.color.rgb = DARK_TEAL

    doc.add_paragraph()
    rule = doc.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = rule.add_run("—" * 20)
    run.font.color.rgb = COPPER

    for _ in range(6):
        doc.add_paragraph()

    for line in meta_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        run.font.name = BODY_FONT
        run.font.size = Pt(10.5)
        run.font.color.rgb = SLATE
    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int = 1):
    return doc.add_heading(text, level=level)


def add_body(doc: Document, text: str, italic: bool = False, bold: bool = False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = italic
    run.bold = bold
    return p


def add_bullets(doc: Document, items: list[str], style: str = "List Bullet") -> None:
    for item in items:
        doc.add_paragraph(item, style=style)


def add_numbered(doc: Document, items: list[str]) -> None:
    add_bullets(doc, items, style="List Number")


def add_callout(doc: Document, label: str, text: str, kind: str = "note") -> None:
    color = {"note": TEAL, "warning": AMBER, "critical": RUST}.get(kind, TEAL)
    p = doc.add_paragraph()
    set_paragraph_shading(p, LIGHT_FILL)
    add_left_border(p, color)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Pt(10)
    label_run = p.add_run(f"{label}  ")
    label_run.bold = True
    label_run.font.color.rgb = color
    label_run.font.name = HEADING_FONT
    label_run.font.size = Pt(10.5)
    body_run = p.add_run(text)
    body_run.font.size = Pt(10.5)
    body_run.font.color.rgb = DARK_TEAL


def add_code_block(doc: Document, lines: list[str]) -> None:
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        set_paragraph_shading(p, CODE_FILL)
        p.paragraph_format.space_before = Pt(0 if i else 4)
        p.paragraph_format.space_after = Pt(4 if i == len(lines) - 1 else 0)
        p.paragraph_format.left_indent = Pt(10)
        run = p.add_run(line if line else " ")
        run.font.name = MONO_FONT
        run.font.size = Pt(10)
        run.font.color.rgb = DARK_TEAL


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True

    # Repeat the header row on every page the table spans across.
    trPr = table.rows[0]._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.name = HEADING_FONT
        run.font.size = Pt(10.5)
        run.font.color.rgb = DARK_TEAL
        set_cell_shading(hdr_cells[i], LIGHT_FILL)

    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            run = p.add_run(str(val))
            run.font.size = Pt(10)
            run.font.name = BODY_FONT
            run.font.color.rgb = DARK_TEAL

    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)

    doc.add_paragraph()
    return table


def add_chart(doc: Document, filename: str, caption: str, width_in: float = 6.2) -> None:
    path = CHARTS_DIR / filename
    if not path.exists():
        raise FileNotFoundError(f"Chart not found: {path}")
    doc.add_picture(str(path), width=Inches(width_in))
    last = doc.paragraphs[-1]
    last.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    run.italic = True
    run.font.size = Pt(9.5)
    run.font.color.rgb = SLATE


def add_page_break(doc: Document) -> None:
    doc.add_page_break()
